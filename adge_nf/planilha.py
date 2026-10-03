"""Planilha Excel (.xlsx) do período: aba Geral colorida (com o que foi considerado nos totais), uma aba por tipo de nota
(NFS-e e NF-e, estas separadas), uma só com as NF-e sem ciência (apenas resumo) e uma de canceladas."""
from __future__ import annotations

import datetime as dt
import re

from . import NOME_APP, SITE_ADGE, cfop as cfop_mod, core, totais

VERDE_FUNDO, VERMELHO_FUNDO = "E3F5E9", "FDE7E7"        # paleta A: tons claros, para não poluir
VERDE_TEXTO, VERMELHO_TEXTO = "1F9D5B", "C0392B"
BANNER, CABECALHO_RESUMO, TEXTO_RESUMO = "1F9D5B", "EEF3F0", "1B5E3B"
FORMATO_REAL = '"R$" #,##0.00'

AMARELO_FUNDO, CINZA_FUNDO = "FFF4CC", "F0F0F0"
COLUNAS = [("Tipo", 32), ("Nº", 12), ("Emissão", 13), ("Emitente", 38), ("CNPJ emitente", 20), ("Tomador / destinatário", 38),
           ("CNPJ tomador / destinatário", 20), ("Valor", 16), ("Valor líquido", 16), ("ISS", 14), ("Chave", 54)]
COLUNAS_GERAL = COLUNAS + [("Considerado nos totais", 22)]
NFE_COLUNAS = [("Nº", 12), ("Série", 8), ("Emissão", 13), ("Emitente", 38), ("CNPJ emitente", 20), ("Destinatário", 38),
               ("CNPJ destinatário", 20), ("CFOP", 16), ("Natureza da operação", 34), ("Valor dos produtos", 18), ("Desconto", 14),
               ("Frete", 14), ("IPI", 12), ("ICMS ST", 14), ("ICMS", 14), ("Valor total da nota", 18), ("Chave", 54)]
RESUMO_COLUNAS = [("Emissão", 13), ("Emitente", 38), ("CNPJ emitente", 20), ("Valor", 16), ("Situação", 46), ("Chave", 54)]
ROTULOS = {"servico_prestado": "Serviço prestado", "servico_tomado": "Serviço tomado",
           "paulistana_prestado": "Nota Paulistana prestada", "paulistana_tomado": "Nota Paulistana tomada",
           **{c: v["rotulo"] for c, v in core.CATEGORIAS.items() if v["grupo"] == "nfe"}}
ABAS = {"servico_prestado": "Prestado", "servico_tomado": "Tomado", "paulistana_prestado": "Paulistana Prest.",
        "paulistana_tomado": "Paulistana Tom.", "nfe_saida": "NF-e Venda", "nfe_entrada": "NF-e Compra",
        "nfe_devolucao_venda": "NF-e Dev. Venda", "nfe_devolucao_compra": "NF-e Dev. Compra", "nfe_outras": "NF-e Outras",
        "nfe_resumo": "NF-e sem ciência"}
DINHEIRO = {"Valor", "Valor líquido", "ISS", "Valor dos produtos", "Desconto", "Frete", "IPI", "ICMS ST", "ICMS", "Valor total da nota"}


def _data(d):
    if d.get("processamento"):
        return core.data_do_doc(d)
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", d.get("emissao") or "")
    return dt.date(int(m[1]), int(m[2]), int(m[3])) if m else None


def _linha(a, ativos=None) -> list:
    d = a["doc"]
    linha = [ROTULOS[a["cat"]], _numero(d["numero"]), _data(d), d["emitente_nome"], d["emitente_doc"],
             d["tomador_nome"], d["tomador_doc"], d["valor"], d["valor_liquido"], d["iss"], d["chave"]]
    if ativos is not None:
        linha.append("Sim" if a["cat"] in ativos else "Não")
    return linha


def _cfops(d) -> str:
    return ", ".join(cfop_mod.formatar(c) for c in d.get("cfops") or [d.get("cfop")] if c)


def _linha_nfe(a) -> list:
    d = a["doc"]
    return [_numero(d["numero"]), d.get("serie") or "", _data(d), d["emitente_nome"], d["emitente_doc"], d["tomador_nome"],
            d["tomador_doc"], _cfops(d), d.get("natureza") or "", d.get("valor_produtos", 0.0), d.get("desconto", 0.0),
            d.get("frete", 0.0), d.get("ipi", 0.0), d.get("st", 0.0), d.get("icms", 0.0), d["valor"], d["chave"]]


def _situacao_resumo(d) -> str:
    return ("Ciência enviada, aguardando o XML completo" if d.get("ciencia_enviada")
            else "Sem ciência da operação: a SEFAZ só entregou o resumo")


def _linha_resumo(a) -> list:
    d = a["doc"]
    return [_data(d), d["emitente_nome"], d["emitente_doc"], d["valor"], _situacao_resumo(d), d["chave"]]


def _fundo_da(a) -> str:
    if a["cat"] == "nfe_resumo":
        return AMARELO_FUNDO
    lado = core.CATEGORIAS[a["cat"]]["lado"]
    return VERDE_FUNDO if lado == "receita" else (VERMELHO_FUNDO if lado == "custo" else CINZA_FUNDO)


def _ordenar(itens):
    return sorted(itens, key=lambda a: (_data(a["doc"]) or dt.date.min, str(a["doc"]["numero"]).zfill(12)))


def _numero(n):
    return int(n) if str(n).isdigit() else n


def _cabecalho(ws, linha, estilos, colunas=None):
    Font, Border, Side, Alignment = estilos["Font"], estilos["Border"], estilos["Side"], estilos["Alignment"]
    colunas = colunas or COLUNAS
    for i, (nome, largura) in enumerate(colunas, 1):
        c = ws.cell(linha, i, nome)
        c.font = Font(bold=True)
        c.border = Border(bottom=Side(style="medium", color="808080"))
        c.alignment = Alignment(horizontal="right" if nome in DINHEIRO else "left", indent=1 if nome == "Chave" else 0)
        ws.column_dimensions[chr(64 + i)].width = largura


def _tabela(ws, primeira, itens, estilos, colorir, ativos=None, colunas=None, montar=None):
    """Escreve as linhas. `montar` devolve os valores de uma nota (padrão: layout geral); dinheiro, data e texto seguem o nome da coluna."""
    PatternFill, Alignment = estilos["PatternFill"], estilos["Alignment"]
    colunas = colunas or COLUNAS
    nomes = [n for n, _ in colunas]
    texto = {"CNPJ emitente", "CNPJ tomador / destinatário", "CNPJ destinatário", "Chave", "CFOP"}
    for k, a in enumerate(itens):
        r = primeira + k
        fundo = PatternFill("solid", fgColor=_fundo_da(a)) if colorir else None
        valores = montar(a) if montar else _linha(a, ativos)
        for i, v in enumerate(valores, 1):
            nome = nomes[i - 1]
            c = ws.cell(r, i, v)
            if fundo:
                c.fill = fundo
            if nome == "Emissão":
                c.number_format = "dd/mm/yyyy"
                c.alignment = Alignment(horizontal="left")
            elif nome in DINHEIRO:
                c.number_format = FORMATO_REAL
            elif nome in texto:
                c.number_format = "@"
                c.value = str(v or "")
                if nome == "Chave":
                    c.alignment = Alignment(horizontal="left", indent=1)
    return primeira + len(itens)


def _banner(ws, titulo, colunas, estilos):
    Font, PatternFill, Alignment = estilos["Font"], estilos["PatternFill"], estilos["Alignment"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=colunas)
    c = ws.cell(1, 1, titulo)
    c.font, c.fill = Font(bold=True, color="FFFFFF", size=13), PatternFill("solid", fgColor=BANNER)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 26


def _rodape(ws, linha, estilos):
    c = ws.cell(linha, 1, f"Gerado pelo {NOME_APP} · {SITE_ADGE}")
    c.font = estilos["Font"](color="808080", size=9)
    c.hyperlink = SITE_ADGE


def gerar_xlsx(caminho, empresa: str, ano: int, mes: int, arquivos: list, resumo: dict, canceladas: list = None, ativos=None):
    """`ativos`: categorias consideradas nos totais (o mesmo conjunto marcado na tela). Sem ele, vale o padrão."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    estilos = dict(Alignment=Alignment, Border=Border, Font=Font, PatternFill=PatternFill, Side=Side)
    periodo = f"{core.mes_exibicao(mes)}/{ano}"
    presentes = [c for c in core.ORDEM_CATEGORIAS if c in resumo and ((core.CATEGORIAS[c]["grupo"] == "nfse" and not core.CATEGORIAS[c].get("origem"))
                 or resumo[c]["qtd"] or resumo[c]["canceladas"])]
    if ativos is None:
        ativos = totais.padrao_ativos(presentes)
    ativos = set(ativos)
    ordenadas = _ordenar(arquivos)
    soma = totais.totais(arquivos, ativos)

    wb = Workbook()
    ws = wb.active
    ws.title = "Geral"
    ws.sheet_view.showGridLines = False
    for i, (_, larg) in enumerate(COLUNAS_GERAL, 1):
        ws.column_dimensions[chr(64 + i)].width = larg
    _banner(ws, f"RESUMO · {empresa} · {periodo}", 5, estilos)
    for i, t in enumerate(["Categoria", "Qtd notas", "Valor", "Canceladas (fora)", "Considerado"], 1):
        c = ws.cell(2, i, t)
        c.font, c.fill = Font(bold=True, color=TEXTO_RESUMO), PatternFill("solid", fgColor=CABECALHO_RESUMO)
        c.alignment = Alignment(horizontal="left" if i == 1 else "right", indent=1)
    for k, cat in enumerate(presentes):
        r, dados = 3 + k, resumo[cat]
        lado = core.CATEGORIAS[cat]["lado"]
        fundo, cor = {"receita": (VERDE_FUNDO, VERDE_TEXTO), "custo": (VERMELHO_FUNDO, VERMELHO_TEXTO)}.get(lado, (CINZA_FUNDO, "555555"))
        if cat == "nfe_resumo":
            fundo, cor = AMARELO_FUNDO, "8A6D00"
        rot = ROTULOS[cat] + (" (faturamento)" if cat == "servico_prestado" else "")
        ws.row_dimensions[r].height = 24
        for i, v in enumerate([rot, dados["qtd"], dados["valor"], dados["canceladas"], "Sim" if cat in ativos else "Não"], 1):
            c = ws.cell(r, i, v)
            c.fill, c.font = PatternFill("solid", fgColor=fundo), Font(bold=True, color=cor, size=12 if i == 3 else 11)
            c.alignment = Alignment(horizontal="left" if i == 1 else "right", vertical="center", indent=1)
            if i == 3:
                c.number_format = FORMATO_REAL
    # totais da seleção, ao lado do resumo (colunas G a I)
    for k, (rot, valor, bom) in enumerate([("Faturamento considerado", soma["faturamento"], True),
                                           ("Compras e serviços tomados considerados", soma["compras"], False),
                                           ("Saldo líquido (faturamento − compras)", soma["saldo"], soma["saldo"] >= 0)]):
        r = 2 + k
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        c1, c2 = ws.cell(r, 7, rot), ws.cell(r, 9, valor)
        fundo = CABECALHO_RESUMO if k == 0 else (VERDE_FUNDO if bom else VERMELHO_FUNDO)
        cor = TEXTO_RESUMO if k == 0 else (VERDE_TEXTO if bom else VERMELHO_TEXTO)
        for c in (c1, c2):
            c.fill, c.font = PatternFill("solid", fgColor=fundo), Font(bold=True, color=cor)
        c1.alignment = Alignment(horizontal="left", indent=1)
        c2.number_format, c2.alignment = FORMATO_REAL, Alignment(horizontal="right")
    topo = 3 + len(presentes) + 1
    _cabecalho(ws, topo, estilos, COLUNAS_GERAL)
    fim = _tabela(ws, topo + 1, ordenadas, estilos, colorir=True, ativos=ativos, colunas=COLUNAS_GERAL)
    ws.auto_filter.ref = f"A{topo}:{chr(64 + len(COLUNAS_GERAL))}{max(fim - 1, topo)}"
    _rodape(ws, fim + 1, estilos)

    for cat in presentes:                                # uma aba por tipo, sem cores
        itens = [a for a in ordenadas if a["cat"] == cat]
        dados = resumo[cat]
        grupo = core.CATEGORIAS[cat]["grupo"]
        w = wb.create_sheet(ABAS[cat])
        w.sheet_view.showGridLines = False
        if cat == "nfe_resumo":
            _aba_sem_ciencia(w, itens, dados, empresa, periodo, cat in ativos, estilos)
            continue
        if grupo == "nfse":
            cols, montar, extra = COLUNAS, None, ("Qtd notas", "Valor", "Valor líquido", "ISS", "Canceladas (fora)")
            vals = [dados["qtd"], dados["valor"], dados["liquido"], dados["iss"], dados["canceladas"]]
        else:
            cols, montar, extra = NFE_COLUNAS, _linha_nfe, ("Qtd notas", "Valor", "ICMS", "Canceladas (fora)")
            vals = [dados["qtd"], dados["valor"], dados.get("icms", 0.0), dados["canceladas"]]
        _banner(w, f"RESUMO · {ROTULOS[cat]} · {empresa} · {periodo}", len(extra), estilos)
        for i, t in enumerate(extra, 1):
            c = w.cell(2, i, t)
            c.font, c.fill = Font(bold=True, color=TEXTO_RESUMO), PatternFill("solid", fgColor=CABECALHO_RESUMO)
            c.alignment = Alignment(horizontal="right")
        for i, v in enumerate(vals, 1):
            c = w.cell(3, i, v)
            c.font, c.alignment = Font(bold=True, size=12), Alignment(horizontal="right")
            if extra[i - 1] in ("Valor", "Valor líquido", "ISS", "ICMS"):
                c.number_format = FORMATO_REAL
        _cabecalho(w, 5, estilos, cols)
        f = _tabela(w, 6, itens, estilos, colorir=False, colunas=cols, montar=montar)
        w.auto_filter.ref = f"A5:{chr(64 + len(cols))}{max(f - 1, 5)}"
        _rodape(w, f + 1, estilos)

    wc = wb.create_sheet("Canceladas")
    wc.sheet_view.showGridLines = False
    _banner(wc, f"NOTAS CANCELADAS (fora dos totais) · {empresa} · {periodo}", 5, estilos)
    _cabecalho(wc, 3, estilos)
    lista = _ordenar(canceladas or [])
    f = _tabela(wc, 4, lista, estilos, colorir=False)
    if not lista:
        wc.cell(2, 1, "Nenhuma nota cancelada neste período.").font = Font(italic=True, color="808080")
    _rodape(wc, f + 1, estilos)
    from openpyxl.worksheet.properties import PageSetupProperties
    for w in wb.worksheets:                              # imprime em paisagem, cabendo na largura da página
        w.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        w.page_setup.orientation, w.page_setup.fitToWidth, w.page_setup.fitToHeight = "landscape", 1, 0
    wb.save(str(caminho))


def _aba_sem_ciencia(w, itens, dados, empresa, periodo, considerada, estilos):
    """NF-e que a SEFAZ só entregou como resumo (sem ciência da operação): ficam separadas e marcadas, nunca misturadas às completas."""
    Font, PatternFill, Alignment = estilos["Font"], estilos["PatternFill"], estilos["Alignment"]
    _banner(w, f"NF-e SEM CIÊNCIA DA OPERAÇÃO (só o resumo) · {empresa} · {periodo}", len(RESUMO_COLUNAS), estilos)
    w.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(RESUMO_COLUNAS))
    c = w.cell(2, 1, "Estas notas ainda não tiveram a Ciência da Operação, então a SEFAZ entregou só o resumo (sem CFOP, itens ou impostos). "
                     + ("Estão sendo contadas nos totais como compra, pelo valor total, sem crédito." if considerada
                        else "Não estão sendo contadas nos totais."))
    c.font, c.fill = Font(italic=True, color="8A6D00"), PatternFill("solid", fgColor=AMARELO_FUNDO)
    c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    w.row_dimensions[2].height = 34
    for i, t in enumerate(["Qtd notas", "Valor"], 1):
        h = w.cell(3, i, t)
        h.font, h.fill = Font(bold=True, color=TEXTO_RESUMO), PatternFill("solid", fgColor=CABECALHO_RESUMO)
        h.alignment = Alignment(horizontal="right")
        v = w.cell(4, i, dados["qtd"] if i == 1 else dados["valor"])
        v.font, v.alignment = Font(bold=True, size=12), Alignment(horizontal="right")
        if i == 2:
            v.number_format = FORMATO_REAL
    _cabecalho(w, 6, estilos, RESUMO_COLUNAS)
    f = _tabela(w, 7, itens, estilos, colorir=True, colunas=RESUMO_COLUNAS, montar=_linha_resumo)
    w.auto_filter.ref = f"A6:{chr(64 + len(RESUMO_COLUNAS))}{max(f - 1, 6)}"
    _rodape(w, f + 1, estilos)
