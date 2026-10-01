import pathlib
import sys
import tempfile
import unittest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from adge_nf import core
import test_core
from test_core import CNPJ, NOME

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


@unittest.skipIf(openpyxl is None, "openpyxl não instalado")
class Planilha(unittest.TestCase):
    def gerar(self, tipos=("prestado", "tomado")):
        extras = {}
        arqs, resumo, _ = core.montar_plano(test_core.Plano().docs(), CNPJ, 2026, 9, set(tipos), extras=extras)
        from adge_nf import planilha
        t = tempfile.mkdtemp()
        caminho = pathlib.Path(t) / "x.xlsx"
        planilha.gerar_xlsx(caminho, NOME, 2026, 9, arqs, resumo, extras.get("canceladas", []))
        return openpyxl.load_workbook(caminho), arqs, resumo, extras

    def test_abas_e_resumo(self):
        wb, arqs, resumo, extras = self.gerar()
        self.assertEqual(wb.sheetnames, ["Geral", "Prestado", "Tomado", "Canceladas"])
        g = wb["Geral"]
        self.assertTrue(g["A1"].value.startswith("RESUMO · " + NOME))
        self.assertEqual([g["A2"].value, g["B2"].value, g["C2"].value], ["Categoria", "Qtd notas", "Valor"])
        self.assertEqual(g["A3"].value, "Serviço prestado (faturamento)")
        self.assertEqual(g["B3"].value, 3)
        self.assertAlmostEqual(g["C3"].value, 3505.5)
        self.assertEqual(g["A4"].value, "Serviço tomado")
        self.assertAlmostEqual(g["C4"].value, 300.25)

    def test_geral_colorida_sem_linhas_em_branco_e_cabecalho_em_negrito(self):
        wb, arqs, *_ = self.gerar()
        g = wb["Geral"]
        topo = 6  # banner, cabeçalho do resumo, 2 categorias, 1 linha em branco
        self.assertTrue(g.cell(topo, 1).font.bold)
        self.assertEqual(g.cell(topo, 1).value, "Tipo")
        linhas = [g.cell(r, 1).value for r in range(topo + 1, topo + 1 + len(arqs))]
        self.assertNotIn(None, linhas)                       # nenhuma linha pulada
        for r in range(topo + 1, topo + 1 + len(arqs)):
            prest = g.cell(r, 1).value == "Serviço prestado"
            self.assertEqual(g.cell(r, 1).fill.fgColor.rgb[-6:], "E3F5E9" if prest else "FDE7E7")

    def test_abas_por_tipo_sem_cor(self):
        wb, *_ = self.gerar()
        p = wb["Prestado"]
        self.assertEqual(p["A3"].value, 3)                   # qtd
        self.assertTrue(p.cell(5, 1).font.bold)
        self.assertIn(p.cell(6, 1).fill.fill_type, (None,))
        self.assertEqual(sum(1 for r in range(6, 20) if p.cell(r, 1).value == "Serviço prestado"), 3)
        t = wb["Tomado"]
        self.assertEqual(t["A3"].value, 1)

    def test_canceladas_e_so_um_tipo(self):
        wb, arqs, resumo, extras = self.gerar(tipos=("prestado",))
        self.assertEqual(wb.sheetnames, ["Geral", "Prestado", "Canceladas"])
        self.assertEqual(len(extras["canceladas"]), 1)
        self.assertEqual(wb["Canceladas"]["D4"].value, NOME)

    def test_chave_e_cnpj_como_texto(self):
        wb, arqs, *_ = self.gerar()
        g = wb["Geral"]
        self.assertEqual(g.cell(7, 11).value, arqs[0]["doc"]["chave"])
        self.assertIsInstance(g.cell(7, 5).value, str)


@unittest.skipIf(openpyxl is None, "openpyxl não instalado")
class PlanilhaNFe(unittest.TestCase):
    def gerar(self, ativos=None):
        from adge_nf import nfe, planilha
        from fixtures import doc_nfe, xml_nfe, xml_resnfe
        eu, outro, cli = "11111111000111", "33333333000133", "22222222000122"
        docs = [doc_nfe(1, xml_nfe(eu, cli, [("5102", 1000.0)], num=1, icms=100.0)), doc_nfe(2, xml_nfe(outro, eu, [("5102", 400.0)], num=2)),
                doc_nfe(3, xml_resnfe(outro, 250.0, num=8), "res")]
        arqs, resumo, _ = core.planejar({"tipos": {"prestado": False, "tomado": False}, "nfe": True}, [], eu, 2026, 9, None, {}, docs_nfe=docs)
        caminho = pathlib.Path(tempfile.mkdtemp()) / "n.xlsx"
        planilha.gerar_xlsx(caminho, "EMPRESA", 2026, 9, arqs, resumo, [], ativos=ativos)
        return openpyxl.load_workbook(caminho)

    def test_abas_e_sem_ciencia_separada(self):
        wb = self.gerar()
        self.assertIn("Canceladas", wb.sheetnames)
        nomes = " ".join(wb.sheetnames)
        self.assertIn("NF-e", nomes)
        sem = [w for w in wb.worksheets if "ciência" in w.title.lower() or "ciencia" in w.title.lower()]
        self.assertEqual(len(sem), 1)
        texto = " ".join(str(c.value) for linha in sem[0].iter_rows() for c in linha if c.value)
        self.assertIn("250", texto)                                                  # a nota resumo aparece na aba própria

    def test_geral_marca_o_que_foi_considerado(self):
        g = self.gerar()["Geral"]
        marca = {g.cell(r, 1).value: g.cell(r, 5).value for r in range(3, 12) if g.cell(r, 5).value in ('Sim', 'Não')}
        resumo = [v for k, v in marca.items() if "sem ciência" in str(k).lower()]
        self.assertEqual(resumo, ["Não"])
        self.assertIn("Sim", marca.values())
        g2 = self.gerar({"nfe_saida", "nfe_entrada", "nfe_resumo"})["Geral"]
        marca2 = {g2.cell(r, 1).value: g2.cell(r, 5).value for r in range(3, 12) if g2.cell(r, 5).value in ('Sim', 'Não')}
        self.assertEqual([v for k, v in marca2.items() if "sem ciência" in str(k).lower()], ["Sim"])


if __name__ == "__main__":
    unittest.main()
