"""Teste de fumaça da interface: abre as janelas, simula uma busca com um ADN falso e confere os totais.
Roda no Windows do GitHub Actions (e em Linux com xvfb). Sai com código != 0 se algo falhar."""
import pathlib
import sys
import tempfile
import time
import traceback

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from adge_nf import app as A  # noqa: E402
from adge_nf import busca as B  # noqa: E402
from adge_nf import core, dialogos as D, ui  # noqa: E402
from adge_nf.store import Armazenamento  # noqa: E402
from fixtures import SessaoFalsa, item, xml_nfse  # noqa: E402

CNPJ = "11222333000181"
mensagens = []
ui.erro = lambda *a, **k: mensagens.append(("erro", a))
ui.avisar = lambda *a, **k: mensagens.append(("aviso", a))
ui.perguntar = lambda *a, **k: (mensagens.append(("pergunta", a)) or True)
ui.informar = lambda *a, **k: mensagens.append(("info", a))
ESCOLHA = {"v": "sair"}                       # resposta combinada para o aviso "as notas não foram salvas"
ui.escolher = lambda *a, **k: (mensagens.append(("escolher", a)) or ESCOLHA["v"])
ui.abrir_pasta = lambda *_: None


class CofreFalso:
    d = {}
    def get_password(self, s, u): return self.d.get((s, u))
    def set_password(self, s, u, p): self.d[(s, u)] = p
    def delete_password(self, s, u): self.d.pop((s, u), None)


def esperar(cond, app, segundos=10):
    fim = time.time() + segundos
    while time.time() < fim:
        app.update()
        if cond():
            return True
        time.sleep(0.05)
    return False


def passo(nome):
    print("ok -", nome, flush=True)


def principal():
    with tempfile.TemporaryDirectory() as t:
        t = pathlib.Path(t)
        (t / "cli").mkdir()
        pfx = t / "c.pfx"
        pfx.write_bytes(b"x")
        store = Armazenamento(t / "dados", keyring_mod=CofreFalso())
        store.preferencias["verificar_atualizacao"] = False  # o teste não conversa com o GitHub de verdade
        store.preferencias["boas_vindas_vista"] = True       # o cartão de boas-vindas tem teste próprio (no fim)
        app = A.App(store)
        app.update()
        passo("janela principal abre")

        # --- cadastro de empresa
        core.ler_certificado = lambda *_: {"cnpj": CNPJ, "nome": "EMPRESA TESTE", "valido_ate": __import__("datetime").date(2099, 1, 1)}
        d = D.DialogoEmpresa(app, store)
        d.update()
        d.v_pfx.set(str(pfx)); d.v_senha.set("segredo"); d._validar()
        assert d.v_cnpj.get() == "11.222.333/0001-81", d.v_cnpj.get()
        assert d.v_nome.get() == "EMPRESA TESTE"
        d.v_destino.set(str(t / "cli")); d.cb_est.current(1)  # ano_mes
        d._salvar()
        assert d.salvou, d.l_erro.cget("text")
        assert "segredo" not in (t / "dados" / "empresas.json").read_text(encoding="utf-8")
        app._atualizar_lista(d.emp["id"])
        assert len(app._cartoes) == 1 and app.selecionada_id == d.emp["id"]
        assert store.empresas[0]["cert_validade"] == "2099-01-01"
        passo("cadastro de empresa salva e lista de cartões atualiza")

        # Editar empresa tem a chave da Paulistana, igual a NF-e (começa desligada e fica gravada)
        de = D.DialogoEmpresa(app, store, store.empresas[0])
        de.update()
        assert de.v_paul.get() is False
        de.v_paul.set(True); de._salvar()
        assert de.salvou and store.empresas[0]["paulistana"] is True
        de = D.DialogoEmpresa(app, store, store.empresas[0])
        de.update()
        assert de.v_paul.get() is True
        de.v_paul.set(False); de._salvar()
        assert de.salvou and store.empresas[0]["paulistana"] is False
        passo("Editar empresa: chave da Nota Paulistana liga, grava e desliga")

        # --- busca
        emp = store.empresas[0]
        eu = dict(emit_cnpj=CNPJ, emit_nome="EMPRESA TESTE")
        sessao = SessaoFalsa([(0, [
            item(1, xml_nfse("1" * 50, valor="1000.00", dh="2026-09-01T10:00:00-03:00", num="1", **eu)),
            item(2, xml_nfse("2" * 50, valor="250.50", dh="2026-09-30T10:00:00-03:00", num="2", **eu)),
            item(3, xml_nfse("3" * 50, valor="80.00", dh="2026-09-02T10:00:00-03:00", num="3", toma_cnpj=CNPJ, toma_nome="EMPRESA TESTE")),
        ])])
        b = B.DialogoBusca(app, store, emp, sessao=sessao)
        b.update()
        assert b.v_mes.get() in core.MESES_TELA
        assert b.ano == __import__("datetime").date.today().year
        b._escolher_mes(8); b.ano = 2026; b._atualizar_meses()
        assert "01/09/2026 a 30/09/2026" in b.l_periodo.cget("text") and b.v_mes.get() == "Setembro", b.l_periodo.cget("text")
        assert b.b_meses[8].cget("text") == "Set"
        b._mudar_ano(-1); assert b.ano == 2025; b._mudar_ano(1)
        b._ir_mes_anterior(); b._ir_este_mes(); b._escolher_mes(8); b.ano = 2026; b._atualizar_meses()
        b._buscar()
        assert esperar(lambda: b.resultado is not None, app), f"busca não terminou: {mensagens}"
        r = b.resultado
        assert abs(r["resumo"]["servico_prestado"]["valor"] - 1250.50) < 0.001
        assert len(b.tv.get_children()) == 3
        assert "1.250,50" in b.cards["fat"][1].cget("text"), b.cards["prestado"][1].cget("text")
        passo("busca mostra faturamento e lista de notas")
        # saldo líquido (prestado − tomado): verde se positivo; totais e botões no mesmo bloco
        assert b.l_saldo_v.cget("text") == "R$ 1.170,50", b.l_saldo_v.cget("text")
        assert b.c_saldo.fundo.upper() == "#E3F5E9"
        assert b.c_saldo.winfo_ismapped()
        assert b.b_copiar.master is b.f_acoes_totais and b.b_avancado.master is b.f_acoes_totais
        assert b.f_acoes_totais.master is b.f_bloco and b.b_copiar.winfo_ismapped()
        assert len(b.cards_pos.winfo_children()) == 3
        assert "Saldo líquido" in b._texto_totais() and "1.170,50" in b._texto_totais()
        neg = SessaoFalsa([(0, [
            item(1, xml_nfse("4" * 50, valor="100.00", dh="2026-09-03T10:00:00-03:00", num="4", **eu)),
            item(2, xml_nfse("5" * 50, valor="500.00", dh="2026-09-04T10:00:00-03:00", num="5", toma_cnpj=CNPJ, toma_nome="EMPRESA TESTE")),
        ])])
        bn = B.DialogoBusca(app, store, emp, sessao=neg)
        bn._escolher_mes(8); bn.ano = 2026; bn._atualizar_meses(); bn._buscar()
        assert esperar(lambda: bn.resultado is not None, app)
        assert bn.l_saldo_v.cget("text") == "-R$ 400,00", bn.l_saldo_v.cget("text")
        assert bn.c_saldo.fundo.upper() == "#FDE7E7" and str(bn.l_saldo_v.cget("fg")).upper() == "#C0392B"
        bn.destroy()
        passo("saldo líquido muda de cor (positivo/negativo)")

        # --- consulta de períodos em lote (v1.7.5): vem desmarcada; marcada, cada mês clicado entra na mesma busca
        lote = SessaoFalsa([(0, [
            item(1, xml_nfse("6" * 50, valor="100.00", dh="2026-08-10T10:00:00-03:00", num="6", **eu)),
            item(2, xml_nfse("7" * 50, valor="200.00", dh="2026-09-10T10:00:00-03:00", num="7", **eu)),
        ])])
        (t / "lote").mkdir()
        bl = B.DialogoBusca(app, store, dict(emp, destino=str(t / "lote")), sessao=lote)      # pasta à parte: não mistura com os outros testes
        bl.update()
        assert bl.v_lote.get() is False
        bl.ano = 2026; bl._escolher_mes(8); bl._escolher_mes(7)
        assert bl._periodos() == [(2026, 8)], bl._periodos()            # sem o lote, clicar troca o mês
        bl._escolher_mes(8)
        bl.v_lote.set(True); bl._alternar_lote()
        bl._escolher_mes(7)                                              # com o lote, agosto entra junto com setembro
        assert bl._periodos() == [(2026, 8), (2026, 9)], bl._periodos()
        assert "2 meses em lote" in bl.l_periodo.cget("text"), bl.l_periodo.cget("text")
        bl._buscar()
        assert esperar(lambda: bl.resultado is not None, app), f"busca em lote não terminou: {mensagens}"
        assert bl.lote_periodos == [(2026, 8), (2026, 9)] and bl.resultado["mes"] == 8
        assert abs(bl.resultado["resumo"]["servico_prestado"]["valor"] - 100.0) < 0.001
        assert "300,00" in bl.l_lote_resumo.cget("text"), bl.l_lote_resumo.cget("text")
        bl._mostrar_periodo(2026, 9)
        assert bl.resultado["mes"] == 9 and abs(bl.resultado["resumo"]["servico_prestado"]["valor"] - 200.0) < 0.001
        bl._gravar()                                                     # grava os dois meses (cada um na sua pasta)
        pastas = {p.parent.name for p in (t / "lote").rglob("*.xml")}
        assert any("08" in n for n in pastas) and any("09" in n for n in pastas), pastas
        assert bl.salvo and {h["mes"] for h in store.preferencias["historico"]} == {8, 9}      # cada mês entra no histórico
        store.preferencias["historico"] = []
        bl.v_lote.set(False); bl._alternar_lote()
        assert len(bl._periodos()) == 1
        bl.destroy()
        passo("consulta em lote: marcar vários meses, uma busca, trocar o mês exibido e gravar todos")

        # filtros, ordenação e detalhe da nota
        b.v_filtro.set("custo"); b._preencher_tabela()
        assert len(b.tv.get_children()) == 1
        b.v_filtro.set("receita"); b._preencher_tabela()
        assert len(b.tv.get_children()) == 2
        b.v_filtro.set("todas"); b.v_texto.set("250,50"); b.update()
        assert len(b.tv.get_children()) == 1, b.tv.get_children()
        b.v_texto.set("")
        b._ordenar("valor"); primeira = b.tv.item(b.tv.get_children()[0], "values")[4]
        assert primeira == "R$ 80,00", primeira
        b._ordenar("valor"); assert b.tv.item(b.tv.get_children()[0], "values")[4] == "R$ 1.000,00"
        b.tv.selection_set(b.tv.get_children()[0]); b.update()
        assert "Nota" in b.l_detalhe.cget("text") and "Valor" in b.l_detalhe.cget("text"), b.l_detalhe.cget("text")
        passo("filtros, ordenação e detalhe da nota")
        assert store.preferencias.get("historico", []) == [], "consulta sem salvar não pode entrar no histórico"

        b._copiar()
        assert "1.250,50" in app.clipboard_get()
        try:
            import openpyxl
        except ImportError:
            openpyxl = None
        if openpyxl:
            alvo = t / "export.xlsx"
            B.filedialog.asksaveasfilename = lambda **k: str(alvo)
            b._exportar_planilha()
            wbk = openpyxl.load_workbook(alvo)
            assert wbk.sheetnames == ["Geral", "Prestado", "Tomado", "Canceladas"], wbk.sheetnames
            assert abs(wbk["Geral"]["C3"].value - 1250.50) < 0.001
            passo("exportar planilha Excel")
        assert b._tem_nao_salvo()
        b._gravar()
        pasta = t / "cli" / "2026" / "09-Setembro"
        nomes = [p.name for p in pasta.iterdir()]
        assert len(nomes) == 4, nomes  # 3 XMLs + relatório
        h = store.preferencias["historico"]
        assert len(h) == 1 and h[0]["mes"] == 9 and h[0]["pasta"] == str(pasta), h
        assert abs(h[0]["prestado"] - 1250.50) < 0.001 and abs(h[0]["tomado"] - 80) < 0.001, h
        assert not b._tem_nao_salvo()
        passo("gravação cria pasta ano/mês e arquivos e só então entra no histórico")

        # --- só calcular não exige destino
        emp2 = dict(emp, acao="calcular", destino="")
        b2 = B.DialogoBusca(app, store, emp2, sessao=sessao)
        b2._escolher_mes(8); b2.ano = 2026; b2._atualizar_meses()
        b2._buscar()
        assert esperar(lambda: b2.resultado is not None, app)
        assert not b2.b_gravar.winfo_ismapped() and b2.b_pasta.winfo_ismapped()
        assert "ainda não foram salvas" in b2.l_destino.cget("text"), b2.l_destino.cget("text")
        ESCOLHA["v"] = "voltar"
        b2._fechar()
        assert b2.winfo_exists() and mensagens[-1][0] == "escolher"      # até no "só o total": avisa que nada foi salvo
        (t / "calc").mkdir()
        B.filedialog.askdirectory = lambda **k: str(t / "calc")
        ESCOLHA["v"] = "salvar"
        b2._fechar()                                                      # "Salvar numa pasta": escolhe, grava e fecha
        assert not b2.winfo_exists() and len(list((t / "calc").iterdir())) == 4
        assert store.preferencias["historico"][0]["pasta"] == str(t / "calc")
        passo("modo só calcular: pode salvar numa pasta ao sair")

        # --- informações avançadas: gráfico + comparativo de regimes
        from adge_nf import avancado as AV
        AV.DialogoAvancado._fluxo_dados = lambda self: None          # sem janelas bloqueantes no teste
        assert b.b_avancado.winfo_ismapped()
        b._avancado()
        av = [w for w in b.winfo_toplevel().winfo_children() if isinstance(w, AV.DialogoAvancado)][0]
        av.update()
        assert [x[0] for x in av.itens][0].startswith("CLIENTE") or av.itens, av.itens
        av.v_modo.set("tipo"); av._mudou()
        assert {x[0] for x in av.itens} == {"Serviços prestados", "Serviços tomados"}, av.itens
        av._clicar(0)
        assert "R$" in av.l_detalhe.cget("text")
        # sem dados fiscais: pede para informar
        assert any("Informar dados fiscais" in str(w.cget("text")) for w in av.f_comp.winfo_children() if isinstance(w, ui.Botao))
        d = AV.DialogoDadosFiscais(av, store, av.emp)
        d.v_rbt.set("360.000,00"); d.v_folha.set("120.000,00"); d.v_regime.set("Lucro Presumido"); d.v_iss.set("5")
        d._salvar()
        assert d.salvou, d.l_erro.cget("text")
        assert store.obter(av.emp["id"])["fiscal"]["regime"] == "presumido"
        av._render_comparativo()
        av.update()
        txt = []
        def varre(w):
            try:
                txt.append(str(w.cget("text")))
            except tk.TclError:
                pass
            for f in w.winfo_children():
                varre(f)
        import tkinter as tk
        varre(av.f_comp)
        junto = " ".join(txt)
        assert "Lucro Presumido" in junto and "Lucro Real" in junto and "IBS/CBS" in junto, junto[:400]
        assert "Simples Nacional" not in junto, "empresa do Presumido não deve ver análise do Simples"
        assert "Mais econômico" in junto and "Créditos considerados" in junto
        passo("informações avançadas: pizza, dados fiscais e comparativo")
        av.destroy()

        # --- pasta fora do padrão Adge: não trava, deixa escolher outra pasta
        (t / "cliente_sem_padrao").mkdir()
        emp3 = dict(emp, estrutura="adge", destino=str(t / "cliente_sem_padrao"))
        b3 = B.DialogoBusca(app, store, emp3, sessao=sessao)
        b3._escolher_mes(8); b3.ano = 2026; b3._atualizar_meses()
        b3._buscar()
        assert esperar(lambda: b3.resultado is not None, app)
        assert str(b3.b_gravar.cget("state")) == "disabled"
        assert "Escolher pasta" in b3.l_destino.cget("text"), b3.l_destino.cget("text")
        assert b3.b_pasta.winfo_ismapped()
        (t / "livre").mkdir()
        B.filedialog.askdirectory = lambda **k: str(t / "livre")
        b3._escolher_pasta()
        assert str(b3.b_gravar.cget("state")) == "normal"
        b3._gravar()
        assert len(list((t / "livre").iterdir())) == 4, list((t / "livre").iterdir())
        b3.destroy()
        passo("pasta fora do padrão Adge permite escolher outra pasta")

        # --- NF-e: busca junto com a NFS-e, checkmarks, contador da SEFAZ e histórico
        import tkinter as tk
        from adge_nf import nfe as NFE
        from adge_nf import nfe_cache
        from fixtures import (SessaoSefazFalsa, chave_nfe, soap_dist, xml_nfe, xml_resnfe)
        OUTRO, CLI = "33333333000133", "44444444000144"

        hist_base = t / "hist"
        sefaz = SessaoSefazFalsa([soap_dist("138", [
            (1, "procNFe_v4.00.xsd", xml_nfe(CNPJ, CLI, [("5102", 2000.0)], num=1, icms=340.0)),
            (2, "procNFe_v4.00.xsd", xml_nfe(OUTRO, CNPJ, [("5102", 600.0)], num=2, icms=100.0)),
            (3, "procNFe_v4.00.xsd", xml_nfe(CNPJ, CLI, [("5949", 70.0)], num=3)),
            (4, "resNFe_v1.01.xsd", xml_resnfe(OUTRO, 250.0, num=9))], 4, 4)])
        emp_n = dict(emp, nfe=True, nfe_ciencia=False)
        bn = B.DialogoBusca(app, store, emp_n, sessao=sessao, sessao_nfe=sefaz, base_historico=str(hist_base))
        bn._escolher_mes(8); bn.ano = 2026; bn._atualizar_meses(); bn._buscar()
        assert esperar(lambda: bn.resultado is not None, app), f"busca com NF-e não terminou: {mensagens}"
        assert len(sefaz.enviados) == 1 and "210210" not in sefaz.enviados[0][1]      # sem ciência: só leu
        assert bn.cards["fat"][1].cget("text") == "R$ 3.250,50", bn.cards["fat"][1].cget("text")
        assert bn.cards["comp"][1].cget("text") == "R$ 680,00", bn.cards["comp"][1].cget("text")
        assert bn.l_saldo_v.cget("text") == "R$ 2.570,50", bn.l_saldo_v.cget("text")
        assert {"nfe_saida", "nfe_entrada", "nfe_outras", "nfe_resumo"} <= set(bn.marcas), set(bn.marcas)
        assert not bn.sel.ativa("nfe_resumo") and not bn.sel.ativa("nfe_outras") and bn.sel.ativa("nfe_saida")
        assert bn.cards["fat"][4].cget("text") == "Faturamento"
        passo("NF-e junto com a NFS-e: cartões somam e notas sem ciência/outras operações começam fora")

        # as marcas mudam cartões, tabela e planilha; resumo conta como compra pelo valor total
        n_antes = len(bn.tv.get_children())
        bn.sel.definir("nfe_resumo", True)
        assert bn.cards["comp"][1].cget("text") == "R$ 930,00", bn.cards["comp"][1].cget("text")
        assert len(bn.tv.get_children()) == n_antes + 1
        bn.sel.definir("nfe_outras", True)
        assert bn.cards["comp"][1].cget("text") == "R$ 930,00"                       # outras operações não somam
        bn.sel.definir("nfe_saida", False)
        assert bn.cards["fat"][1].cget("text") == "R$ 1.250,50"
        assert "NF-e sem ciência" in bn._texto_totais() or "Considerado" in bn._texto_totais()
        linhas = bn.tv.get_children()
        assert any("semciencia" in bn.tv.item(i, "tags") for i in linhas)
        bn.sel.definir("nfe_saida", True); bn.sel.definir("nfe_resumo", False); bn.sel.definir("nfe_outras", False)
        assert bn.cards["fat"][1].cget("text") == "R$ 3.250,50"
        passo("checkmarks atualizam cartões, tabela e texto copiado")

        # informações avançadas escutam a mesma seleção
        bn._avancado()
        av = [w for w in bn.winfo_toplevel().winfo_children() if isinstance(w, AV.DialogoAvancado)][0]
        av.update()
        assert abs(sum(v for _, v, _ in av.itens) - (3250.50 - 680.0)) < 0.01, av.itens          # o gráfico é líquido: compras reduzem
        assert any(v < 0 for _, v, _ in av.itens) and "Entradas" in av.l_balanco.cget("text")
        bn.sel.definir("nfe_entrada", False); bn.sel.definir("servico_tomado", False)
        av.update()
        assert abs(sum(v for _, v, _ in av.itens) - 3250.50) < 0.01, av.itens
        av.marcas["nfe_entrada"].set(True); av.sel.definir("nfe_entrada", True)
        assert bn.sel.ativa("nfe_entrada") and bn.marcas["nfe_entrada"].get()
        for c in list(bn.sel.ativos):
            bn.sel.definir(c, False)
        av.update()
        assert av.itens == [] and av.l_balanco.cget("text") == ""                                 # nada marcado: zero
        for c in ("servico_prestado", "nfe_saida", "nfe_entrada", "servico_tomado"):
            bn.sel.definir(c, True)
        av.update()
        assert abs(sum(v for _, v, _ in av.itens) - (3250.50 - 680.0)) < 0.01
        av.v_modo.set("servico"); av._mudou()
        assert any(x[0].startswith("CFOP 5.102") for x in av.itens), av.itens
        av.destroy()
        bn.sel.definir("servico_tomado", True)
        passo("informações avançadas usam a mesma seleção e mostram CFOP")

        # a consulta bloqueou a SEFAZ, mas como nada foi salvo o NSU não andou e nenhum XML entrou no controle
        h = nfe_cache.HistoricoNFe(emp_n, base=str(hist_base))
        assert h.bloqueado() and h.estado["ult_nsu"] == 0 and h.quantidade() == 0, h.estado
        bn.update(); bn._tic_nfe()
        assert bn.sw_nfe._desab and not bn.v_nfe.get(), "NF-e deve ficar desligada e travada durante o bloqueio"
        assert "NF-e desligada" in bn.l_nfe_timer.cget("text"), bn.l_nfe_timer.cget("text")
        passo("SEFAZ bloqueada: NF-e desliga e trava; NSU não anda sem salvar")

        # sair sem salvar: aviso (com a nota sobre a SEFAZ); "Voltar" mantém a janela
        ESCOLHA["v"] = "voltar"
        bn._fechar()
        assert bn.winfo_exists() and mensagens[-1][0] == "escolher" and "SEFAZ já bloqueou" in mensagens[-1][1][2], mensagens[-1]
        # salvar: o histórico nasce, com a pasta, e só então o NSU avança e os XMLs entram no controle
        bn._gravar()
        h = nfe_cache.HistoricoNFe(emp_n, base=str(hist_base))
        assert h.estado["ult_nsu"] == 4 and h.quantidade() == 4, h.estado
        assert store.preferencias["historico"][0]["pasta"] == str(t / "cli" / "2026" / "09-Setembro")
        passo("salvar cria o histórico (com a pasta) e confirma o NSU da NF-e")

        # com a SEFAZ bloqueada a próxima busca não toca nela: só NFS-e
        bn.resultado = None
        bn._buscar()
        assert esperar(lambda: bn.resultado is not None, app)
        assert len(sefaz.enviados) == 1, "não deveria consultar a SEFAZ enquanto bloqueada"
        assert bn.cards["fat"][1].cget("text") == "R$ 1.250,50", bn.cards["fat"][1].cget("text")
        # prazo vencido: a NF-e volta sozinha, consulta a partir do NSU confirmado e reaproveita as notas já guardadas
        e_ = h.estado; e_["bloqueio_ate"] = "2000-01-01T00:00:00"; h.salvar_estado()
        bn._tic_nfe()
        assert not bn.sw_nfe._desab and bn.v_nfe.get()
        bn.resultado = None
        bn._buscar()
        assert esperar(lambda: bn.resultado is not None, app)
        assert len(sefaz.enviados) == 2 and "000000000000004" in sefaz.enviados[1][1], sefaz.enviados[1][1]
        assert bn.cards["fat"][1].cget("text") == "R$ 3.250,50", bn.cards["fat"][1].cget("text")
        bn._tic_nfe()
        assert bn.sw_nfe._desab, "a SEFAZ voltou a bloquear (137): NF-e trava de novo"
        passo("NF-e volta quando o prazo acaba, parte do NSU confirmado e reaproveita o que já foi guardado")

        # fechar sem salvar escolhendo "Sair sem salvar": nada entra no histórico novo
        outro = B.DialogoBusca(app, store, dict(emp, nfe=False), sessao=sessao)
        outro._escolher_mes(7); outro.ano = 2026; outro._atualizar_meses(); outro._buscar()
        assert esperar(lambda: outro.resultado is not None, app)
        n_hist = len(store.preferencias["historico"])
        ESCOLHA["v"] = "sair"
        outro._fechar()
        assert not outro.winfo_exists() and len(store.preferencias["historico"]) == n_hist
        # "Salvar numa pasta" a partir do aviso grava, registra e fecha
        ESCOLHA["v"] = "salvar"
        bn._fechar()
        assert not bn.winfo_exists()
        passo("aviso ao sair: sair sem salvar não cria histórico; salvar grava e fecha")

        # --- Nota Paulistana: busca separada, fica desmarcada, e a conferência aparece nos avisos
        from adge_nf import nfe as NFE2
        from fixtures import assinador_falso, soap_sp, xml_nfe_sp
        class SessaoSP:
            def __init__(self, r): self.r, self.enviados = list(r), []
            def post(self, url, data=None, headers=None, timeout=0):
                self.enviados.append(url)
                from fixtures import RespSoap
                return RespSoap(self.r.pop(0))
        sp = SessaoSP([soap_sp("ConsultaNFeRecebidas", []), soap_sp("ConsultaCNPJ", detalhes=["999"]),
                       soap_sp("ConsultaNFeEmitidas", [xml_nfe_sp("999", 77, "2026-09-20", 1234.0, CNPJ, CLI)])])
        NFE2.carregar_assinador = lambda pfx, senha: assinador_falso()
        bp = B.DialogoBusca(app, store, dict(emp, tipos={"prestado": True, "tomado": True}, paulistana=True), sessao=sessao, sessao_paulistana=sp)
        assert bp.v_paul.get() is True               # vem ligada porque a empresa tem a Paulistana ligada no cadastro
        bp._escolher_mes(8); bp.ano = 2026; bp._atualizar_meses(); bp._buscar()
        assert esperar(lambda: bp.resultado is not None, app), f"busca com Paulistana não terminou: {mensagens}"
        assert len(sp.enviados) == 3
        assert "paulistana_prestado" in bp.marcas and not bp.sel.ativa("paulistana_prestado")
        txt = " ".join(bp.resultado["avisos"])
        assert "Só na Paulistana (1): nº 77" in txt and "Só no Ambiente Nacional" in txt, txt
        bp.sel.definir("paulistana_prestado", True)
        bp.update()
        bp.destroy()
        passo("Nota Paulistana: busca separada com conferência")

        # --- erro do ADN aparece sem travar
        class Ruim:
            def get(self, *a, **k):
                return type("R", (), {"status_code": 403, "json": lambda s: {}})()
        b3 = B.DialogoBusca(app, store, emp, sessao=Ruim())
        b3._buscar()
        assert esperar(lambda: any(m[0] == "erro" for m in mensagens), app)
        assert not b3.trabalhando
        b3.destroy()
        passo("erro do ADN vira mensagem")

        # --- atualização: aviso, baixar com integridade, instalar, lembrar depois, pular
        import hashlib
        from adge_nf import atualizacao
        from adge_nf import REPO_GITHUB
        conteudo = b"msi-de-teste" * 5000
        sha = hashlib.sha256(conteudo).hexdigest()
        url = f"https://github.com/{REPO_GITHUB}/releases/download/v9.9.9/App.msi"

        class RespA:
            def __init__(self, code=200, j=None, c=b""):
                self.status_code, self._j, self._c, self.headers = code, j, c, {"Content-Length": str(len(c))}
            def json(self): return self._j
            def iter_content(self, n):
                for i in range(0, len(self._c), n):
                    yield self._c[i:i + n]

        class SessaoA:
            def __init__(self, c=conteudo): self.c = c
            def get(self, u, **k):
                if u.endswith("/releases/latest"):
                    return RespA(200, {"tag_name": "v9.9.9", "html_url": "https://x", "body": "Agora suporta NF-e",
                                       "assets": [{"name": "App.msi", "browser_download_url": url, "size": len(conteudo), "digest": "sha256:" + sha}]})
                return RespA(200, c=self.c)

        app.sessao_att = SessaoA()
        app.pausa_log_att = 0
        abertos = []
        orig = A.DialogoAtualizacao
        def fabrica(pai, st, info, **k):
            d = orig(pai, st, info, ao_instalar=lambda caminho: abertos.append(caminho), pode_instalar=True, sessao=app.sessao_att)
            d.after(100, d._atualizar)           # simula o clique em "Atualizar agora"
            d.after(2500, lambda: d.winfo_exists() and d.destroy())
            return d
        A.DialogoAtualizacao = fabrica
        app._checar_atualizacao(manual=True)
        assert esperar(lambda: abertos, app, 15), f"instalação não foi disparada: {mensagens}"
        assert pathlib.Path(abertos[0]).read_bytes() == conteudo
        assert "Atualização necessária" in app.l_att.cget("text") and "Pronto" in app.l_att.cget("text"), app.l_att.cget("text")
        passo("atualização baixa, confere a integridade e dispara a instalação")
        A.DialogoAtualizacao = orig

        # arquivo adulterado: recusa e não instala
        abertos.clear()
        app.sessao_att = SessaoA(conteudo[:-1] + b"X")
        app.info_att = None
        A.DialogoAtualizacao = fabrica
        app._checar_atualizacao(manual=True)
        esperar(lambda: False, app, 4)
        assert not abertos
        A.DialogoAtualizacao = orig
        passo("arquivo adulterado não é instalado")

        # lembrar depois e pular
        info = atualizacao.consultar(SessaoA(), versao="1.0.0")
        d = orig(app, store, info, pode_instalar=False)
        d.update()
        assert "Agora suporta NF-e" in d.txt.get("1.0", "end")
        assert d.b_ok.cget("text") == "Abrir a página de download"
        d._depois()
        assert not atualizacao.deve_avisar(info, store.preferencias)
        store.preferencias["atualizacao"].pop("lembrar_ate")
        d = orig(app, store, info, pode_instalar=True)
        d.update()
        assert d.b_ok.cget("text") == "Atualizar agora"
        d._pular()
        assert store.preferencias["atualizacao"]["pular"] == "v9.9.9"
        passo("lembrar depois e pular versão")

        # novidades por item com código, versões puladas e link do relatório completo
        corpo = "## Resumo\n- [NOVO] Informações avançadas (ITEM-10): gráfico novo.\n- [CORREÇÃO] Saldo líquido (ITEM-05): cor certa.\n"
        class SessaoB(SessaoA):
            def get(self, u, **k):
                if u.endswith("/releases"):
                    return RespA(200, [{"tag_name": "v9.9.9", "body": corpo},
                                       {"tag_name": "v9.9.8", "body": "## Resumo\n- [MELHORIA] Planilha Excel (ITEM-09): abas.\n"}])
                r = super().get(u, **k)
                if u.endswith("/releases/latest"):
                    r._j["body"] = corpo
                return r
        info = atualizacao.consultar(SessaoB(), versao="1.0.0")
        d = orig(app, store, info, pode_instalar=False)
        d.update()
        texto = d.txt.get("1.0", "end")
        assert "Informações avançadas" in texto and "(ITEM-10)" in texto and "CORREÇÃO" in texto.upper(), texto
        assert texto.index("Versão 9.9.9") < texto.index("Versão 9.9.8") and "(ITEM-09)" in texto
        assert d.l_relatorio.cget("text").startswith("Ver relatório completo")
        d.destroy()
        passo("aviso de atualização lista as novidades por item e versão")

        # --- boas-vindas: só na primeira abertura; "Sobre" permite rever
        assert A.boas_vindas_pendente({}) and not A.boas_vindas_pendente({"boas_vindas_vista": True})
        store.preferencias.pop("boas_vindas_vista")
        abertos_url = []
        ui.abrir_link = lambda u: abertos_url.append(u)
        d = A.DialogoBoasVindas(app, store)
        d.update()
        assert "sem fins lucrativos" in " ".join(t for _, t in A.TEXTOS_MISSAO)
        assert [b.cget("text") for b in (d.b_github, d.b_google, d.b_site)] == ["Avaliar no GitHub", "Avaliar no Google", "Conhecer a Adge"]
        d.b_github.invoke(); d.b_google.invoke()
        assert any("github.com" in u for u in abertos_url) and any("share.google/d9zzh3PZHLl5LCvF2" in u for u in abertos_url), abertos_url
        assert d.b_primeira is None  # já existe empresa cadastrada
        d.b_ok.invoke()
        assert store.preferencias["boas_vindas_vista"] is True and not A.boas_vindas_pendente(store.preferencias)
        app._ir("sobre"); app.update()
        assert app.b_apresentacao.cget("text") == "Ver a apresentação do projeto"
        # primeira abertura sem empresas: oferece cadastrar a primeira
        vazio = Armazenamento(t / "dados2", keyring_mod=CofreFalso())
        chamou = []
        d = A.DialogoBoasVindas(app, vazio, ao_adicionar=lambda: chamou.append(1))
        d.update()
        assert d.b_primeira is not None
        d.b_primeira.invoke()
        assert chamou and vazio.preferencias["boas_vindas_vista"] is True
        passo("boas-vindas aparece uma vez, oferece a primeira empresa e Sobre permite rever")

        # --- histórico, tema escuro e validade do certificado
        app._ir("historico"); app.update()
        assert len(app.rol_hist.corpo.winfo_children()) == 1
        abertas, dialogos = [], []
        orig_busca = A.DialogoBusca

        class Proibida:
            def get(self, *a, **k): raise AssertionError("reabrir pelo histórico não pode consultar o ADN")
            def post(self, *a, **k): raise AssertionError("reabrir pelo histórico não pode consultar a SEFAZ nem a Prefeitura")

        def fabrica_busca(pai, st, e, periodo=None, auto=False, da_pasta=None):
            abertas.append((periodo, auto, da_pasta))
            fontes = dict(sessao_nfe=Proibida(), sessao_paulistana=Proibida()) if da_pasta else {}
            dlg = orig_busca(pai, st, e, sessao=Proibida() if da_pasta else sessao, periodo=periodo, auto=auto, da_pasta=da_pasta, **fontes)
            dialogos.append(dlg)
            def fim():
                if dlg.winfo_exists():
                    dlg.destroy() if dlg.resultado else dlg.after(100, fim)
            dlg.after(200, fim)
            return dlg
        A.DialogoBusca = fabrica_busca
        reg = store.preferencias["historico"][0]
        pasta_h = str(t / "cli" / "2026" / "09-Setembro")
        assert reg["pasta"] == pasta_h
        app._abrir_historico(reg)
        assert abertas == [((2026, 9), False, pasta_h)], abertas        # abre pela pasta, sem busca automática
        r = dialogos[0].resultado
        assert abs(r["resumo"]["servico_prestado"]["valor"] - 1250.50) < 0.001 and abs(r["resumo"]["nfe_saida"]["valor"] - 2000.0) < 0.001, r["resumo"]
        assert dialogos[0].da_pasta == pasta_h and dialogos[0].salvo
        # entrada antiga (sem pasta) ou pasta que sumiu: oferece uma consulta nova
        app._abrir_historico(dict(reg, pasta=None))
        app._abrir_historico(dict(reg, pasta=str(t / "sumiu")))
        assert abertas[1:] == [((2026, 9), True, None), ((2026, 9), True, None)], abertas
        A.DialogoBusca = orig_busca
        n = len(mensagens)
        app._abrir_historico({"empresa_id": "nao-existe", "ano": 2026, "mes": 1})
        assert mensagens[-1][0] == "aviso" and len(mensagens) == n + 1
        app._limpar_historico()
        assert store.preferencias["historico"] == []
        passo("histórico lista, reabre a consulta e limpa")

        assert A.situacao_certificado("2099-01-01")[1] == "verde"
        assert A.situacao_certificado("2000-01-01")[1] == "vermelho"
        assert A.situacao_certificado((__import__("datetime").date.today() + __import__("datetime").timedelta(days=10)).isoformat())[1] == "amarelo"
        assert A.situacao_certificado(None) is None
        store.empresas[0].pop("cert_validade", None); store.empresas[0].pop("cert_checado", None)
        app._atualizar_validades()
        assert esperar(lambda: store.empresas[0].get("cert_validade") == "2099-01-01", app, 6), store.empresas[0]
        passo("validade do certificado é atualizada em segundo plano")

        app._ir("config"); app.update()
        app.v_escuro.set(True); app._alternar_tema()
        assert esperar(lambda: ui.P.nome == "escuro" and app.pagina == "config" and str(app.lateral.cget("bg")).upper() == "#0A3822", app), ui.P.nome
        assert store.preferencias["tema"] == "escuro"
        app._ir("empresas"); app.update()
        app.v_escuro.set(False) if hasattr(app, "v_escuro") else None
        app._ir("config"); app.v_escuro.set(False); app._alternar_tema()
        assert esperar(lambda: ui.P.nome == "claro" and app.pagina == "config", app)
        passo("modo escuro liga, desliga e guarda a preferência")

        # --- configurações
        store.trocar_modo("mestra123")
        app._atualizar_seguranca()
        assert "senha mestra" in app.l_seg.cget("text").lower()
        passo("configurações")
        # a aba de configurações tem rolagem e botão de apagar dados
        assert any(isinstance(w, ui.Rolavel) for w in app.paginas["config"].winfo_children())
        app._apagar_dados()
        assert store.empresas == [] and not (t / "dados" / "empresas.json").exists()
        passo("apagar todos os dados")  # o app fecha sozinho depois de apagar


try:
    principal()
    print("SMOKE OK")
except Exception:
    traceback.print_exc()
    sys.exit(1)
